import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Use Word's built-in highlight constants
from docx.enum.text import WD_COLOR_INDEX

doc = Document()

# ---- Page setup ----
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 2.0

rPr = style.element.get_or_add_rPr()
rFonts = rPr.makeelement(qn('w:rFonts'), {})
rFonts.set(qn('w:eastAsia'), '宋体')
rPr.insert(0, rFonts)

# ---- Read markdown ----
with open(r'D:\my-paper\drafts\Manuscript_JED_RED_revised.md', 'r', encoding='utf-8') as f:
    text = f.read()

lines = text.split('\n')

HL_COLOR = WD_COLOR_INDEX.YELLOW  # Word built-in yellow highlight


def add_heading_styled(doc, text, level, highlight=False):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
        if highlight:
            run.font.highlight_color = HL_COLOR
    return h


def add_paragraph_styled(doc, text, indent=True, center=False, font_size=Pt(12)):
    """Parse paragraph, apply WORD BUILT-IN YELLOW HIGHLIGHT to revised paragraphs."""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    parts = re.split(r'(\*\*.*?\*\*)', text)
    is_revised = any(part.startswith('**') and part.endswith('**') for part in parts)
    # Exclude purely-formatting bolds
    stripped = text.strip()
    if stripped == '**Keywords**:' or stripped.startswith('**Keywords**'):
        is_revised = False

    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            run = p.add_run(part)
        run.font.name = 'Times New Roman'
        run.font.size = font_size
        if is_revised:
            run.font.highlight_color = HL_COLOR
    return p


def add_table_from_md(doc, table_lines):
    rows = []
    for line in table_lines:
        line = line.strip()
        if not line.startswith('|'):
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if all(re.match(r'^[-:—]+$', c) for c in cells):
            continue
        rows.append(cells)
    if len(rows) < 1:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j < ncols:
                cell = table.rows[i].cells[j]
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


# ---- Parse and build ----
i = 0
in_table = False
table_buffer = []
in_references = False

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped == '## References' or stripped.startswith('## References'):
        in_references = True
        add_heading_styled(doc, 'References', level=2)
        i += 1
        continue

    if not stripped:
        if in_table and table_buffer:
            add_table_from_md(doc, table_buffer)
            table_buffer = []
            in_table = False
        i += 1
        continue

    # Title
    if stripped.startswith('# ') and not stripped.startswith('## '):
        title_text = stripped[2:]
        h = doc.add_heading(title_text, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in h.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.highlight_color = HL_COLOR
            run.font.color.rgb = RGBColor(0, 0, 0)
        i += 1
        continue

    # Section headings
    if stripped.startswith('## ') and not in_references:
        heading_text = stripped[3:]
        has_bold = '**' in heading_text
        clean_heading = heading_text.replace('**', '')
        add_heading_styled(doc, clean_heading, level=2, highlight=has_bold)
        i += 1
        continue

    if stripped.startswith('### '):
        heading_text = stripped[4:]
        has_bold = '**' in heading_text
        clean_heading = heading_text.replace('**', '')
        add_heading_styled(doc, clean_heading, level=3, highlight=has_bold)
        i += 1
        continue

    # Metadata
    if stripped.startswith('Min Qu,') or stripped.startswith('Contact information') or stripped.startswith('*Corresponding'):
        add_paragraph_styled(doc, stripped, indent=False, center=True)
        i += 1
        continue

    if stripped == '---':
        i += 1
        continue

    if stripped == '## Abstract' or stripped.startswith('## Abstract'):
        add_heading_styled(doc, 'Abstract', level=2)
        i += 1
        continue

    # Keywords
    if stripped.startswith('**Keywords'):
        add_paragraph_styled(doc, stripped, indent=False)
        i += 1
        continue

    # Tables
    if stripped.startswith('|'):
        in_table = True
        table_buffer.append(stripped)
        i += 1
        continue

    if in_table:
        add_table_from_md(doc, table_buffer)
        table_buffer = []
        in_table = False

    # Figure / table captions
    if stripped.startswith('Fig.') or stripped.startswith('#### Table'):
        add_paragraph_styled(doc, stripped, indent=False, font_size=Pt(10))
        if doc.paragraphs:
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue

    # Math
    if stripped.startswith('$$') or stripped.startswith('\\['):
        add_paragraph_styled(doc, stripped, indent=False)
        if doc.paragraphs:
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue

    # Body paragraph
    add_paragraph_styled(doc, stripped, indent=True)
    i += 1

# Save
output_path = r'D:\my-paper\drafts\Manuscript_JED_RED_revised.docx'
doc.save(output_path)
print(f'Saved to: {output_path}')
